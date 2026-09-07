"""读取上传的 CSV / Excel 文件，返回 [{表头: 值}, ...] 供各工具的导入接口使用。"""
import csv
import io

from fastapi import HTTPException, UploadFile


def read_rows(file: UploadFile) -> list[dict]:
    name = (file.filename or "").lower()
    if name.endswith(".csv"):
        text = file.file.read().decode("utf-8-sig")
        return [row for row in csv.DictReader(io.StringIO(text)) if any(row.values())]

    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(file.file, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        result = []
        for values in rows:
            row = {header[i]: v for i, v in enumerate(values) if i < len(header)}
            if any(v not in (None, "") for v in row.values()):
                result.append(row)
        wb.close()
        return result

    raise HTTPException(400, "仅支持 .csv / .xlsx 文件")
